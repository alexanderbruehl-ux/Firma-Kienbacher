"""Baut aus beliebigen Dateien eine eigenstaendige HTML-Uebersicht mit
Download-Button je Datei.

    python vorschau_erzeugen.py datei1.pptx datei2.png ... --out vorschau.html

Kein Renderversuch ueber eine externe Office-Suite (LibreOffice haengt sich
in vielen Sandbox-Umgebungen ohne Fehlermeldung auf, auch bei trivialen
Dateien - und ein eingebettetes docProps/thumbnail.jpeg ist meist nur ein
leerer Platzhalter, kein echtes Abbild). Stattdessen:

- Bilder und PDFs werden inline angezeigt, weil Browser das nativ koennen.
- PPTX: die erste Folie wird aus den tatsaechlichen Shapes (Position, Farbe,
  Text) als kleine Facsimile-Vorschau nachgebaut - echte Werte aus der
  Datei, kein Rendering-Dienst noetig.
- DOCX: ein Textauszug (Ueberschriften/Absaetze) aus dem Dokumentkoerper.
- Alles andere (XLSX, ZIP, ...) bekommt eine Karte mit Dateityp und Groesse.

Bearbeiten, Drucken und die vollstaendige Ansicht passieren im jeweiligen
Originalprogramm, nicht im Browser - die Vorschau ist ein Anhaltspunkt, kein
Ersatz.

Funktioniert nur als eigenstaendige Datei (lokal geoeffnet oder per Datei
weitergegeben) - in einer eingebetteten/sandboxed Ansicht (z. B. ein
Artifact-Link) sind von der Seite selbst ausgeloeste Downloads blockiert.

Benoetigt `python-pptx` fuer PPTX- und `python-docx` fuer DOCX-Vorschauen;
fehlen sie, faellt die Karte automatisch auf den reinen Dateityp-Hinweis
zurueck.
"""
from __future__ import annotations

import argparse
import base64
import mimetypes
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    _PPTX_VERFUEGBAR = True
except ImportError:
    _PPTX_VERFUEGBAR = False

try:
    import docx as docx_lib
    _DOCX_VERFUEGBAR = True
except ImportError:
    _DOCX_VERFUEGBAR = False

BILD_TYPEN = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".svg"}
PDF_TYPEN = {".pdf"}
EMU_PER_PT = 12700

TYP_LABEL = {
    ".pptx": "PowerPoint", ".ppt": "PowerPoint",
    ".docx": "Word", ".doc": "Word",
    ".xlsx": "Excel", ".xls": "Excel",
    ".pdf": "PDF", ".zip": "ZIP-Archiv",
    ".png": "Bild (PNG)", ".jpg": "Bild (JPEG)", ".jpeg": "Bild (JPEG)",
    ".gif": "Bild (GIF)", ".webp": "Bild (WebP)", ".svg": "Vektorgrafik (SVG)",
    ".json": "JSON", ".md": "Markdown", ".txt": "Text",
}


def esc(t: str) -> str:
    return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _shape_rgb(fill_or_line) -> str | None:
    try:
        return str(fill_or_line.fore_color.rgb) if fill_or_line.type is not None else None
    except Exception:
        return None


def pptx_vorschau_html(pfad: Path) -> str | None:
    """Baut die erste Folie aus den echten Shapes nach - Position, Farbe und
    Text kommen direkt aus der Datei, kein Rendering-Dienst noetig."""
    if not _PPTX_VERFUEGBAR:
        return None
    try:
        prs = Presentation(str(pfad))
        if not prs.slides:
            return None
        slide = prs.slides[0]
        sb_w, sb_h = prs.slide_width, prs.slide_height
        if not sb_w or not sb_h:
            return None
        breite_pt = sb_w / EMU_PER_PT
    except Exception:
        return None

    def pct(v, ganz):
        return round((v or 0) / ganz * 100, 3)

    teile = []
    for shp in slide.shapes:
        links, oben = pct(shp.left, sb_w), pct(shp.top, sb_h)
        breite, hoehe = pct(shp.width, sb_w), pct(shp.height, sb_h)
        basis = f"left:{links}%;top:{oben}%;width:{breite}%;height:{hoehe}%;position:absolute;"

        if shp.shape_type == MSO_SHAPE_TYPE.PICTURE:
            try:
                bild = shp.image
                b64 = base64.b64encode(bild.blob).decode()
                teile.append(f'<img style="{basis}object-fit:contain;" '
                            f'src="data:{bild.content_type};base64,{b64}" alt="" />')
            except Exception:
                pass
            continue

        if shp.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
            try:
                fuellung = _shape_rgb(shp.fill)
            except Exception:
                fuellung = None
            if fuellung:
                teile.append(f'<div style="{basis}background:#{fuellung};"></div>')

        if getattr(shp, "has_text_frame", False) and shp.text_frame.text.strip():
            zeilen = []
            for para in shp.text_frame.paragraphs:
                stueck = []
                for run in para.runs:
                    if not run.text:
                        continue
                    grad_pt = run.font.size.pt if run.font.size else 12
                    grad_cqw = round(grad_pt / breite_pt * 100, 3)
                    fett = "700" if run.font.bold else "400"
                    farbe = _shape_rgb(run.font.color) if run.font.color else None
                    stil = f"font-size:{grad_cqw}cqw;font-weight:{fett};"
                    stil += f"color:#{farbe};" if farbe else ""
                    stueck.append(f'<span style="{stil}">{esc(run.text)}</span>')
                if stueck:
                    zeilen.append("<p style=\"margin:0;line-height:1.2;\">" + "".join(stueck) + "</p>")
            if zeilen:
                teile.append(f'<div style="{basis}overflow:hidden;">{"".join(zeilen)}</div>')

    aspekt = round(sb_w / sb_h, 4)
    inhalt = "\n".join(teile)
    return (f'<div class="preview pptx-preview" style="container-type:inline-size;">'
           f'<div class="pptx-slide" style="aspect-ratio:{aspekt};">{inhalt}</div></div>')


def docx_vorschau_html(pfad: Path) -> str | None:
    """Textauszug (Ueberschriften/Absaetze) aus dem Dokumentkoerper."""
    if not _DOCX_VERFUEGBAR:
        return None
    try:
        dok = docx_lib.Document(str(pfad))
    except Exception:
        return None

    bloecke = []
    zeichen = 0
    for para in dok.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        ist_ueberschrift = para.style is not None and "Heading" in (para.style.name or "")
        tag = "h3" if ist_ueberschrift else "p"
        bloecke.append(f'<{tag}>{esc(text)}</{tag}>')
        zeichen += len(text)
        if zeichen > 700:
            bloecke.append('<p class="fade-hint">&hellip;</p>')
            break
    if not bloecke:
        return None
    return f'<div class="preview docx-preview"><div class="docx-page">{"".join(bloecke)}</div></div>'


def groesse_lesbar(bytes_zahl: int) -> str:
    wert = float(bytes_zahl)
    for einheit in ("B", "KB", "MB", "GB"):
        if wert < 1024 or einheit == "GB":
            return f"{wert:.0f} {einheit}" if einheit == "B" else f"{wert:.1f} {einheit}"
        wert /= 1024
    return f"{wert:.1f} GB"


def datei_karte(pfad: Path) -> str:
    ext = pfad.suffix.lower()
    daten = pfad.read_bytes()
    b64 = base64.b64encode(daten).decode()
    mime = mimetypes.guess_type(pfad.name)[0] or "application/octet-stream"
    typ_label = TYP_LABEL.get(ext, ext.lstrip(".").upper() or "Datei")
    groesse = groesse_lesbar(len(daten))
    data_url = f"data:{mime};base64,{b64}"

    vorschau = None
    if ext in BILD_TYPEN:
        vorschau = f'<div class="preview img-preview"><img src="{data_url}" alt="{esc(pfad.name)}" /></div>'
    elif ext in PDF_TYPEN:
        vorschau = f'<div class="preview pdf-preview"><embed src="{data_url}" type="application/pdf" /></div>'
    elif ext == ".pptx":
        vorschau = pptx_vorschau_html(pfad)
    elif ext == ".docx":
        vorschau = docx_vorschau_html(pfad)

    if vorschau is None:
        vorschau = f'<div class="preview no-preview"><span class="file-badge">{esc(typ_label)}</span></div>'

    return f'''
<article class="card">
  {vorschau}
  <div class="meta">
    <h2>{esc(pfad.name)}</h2>
    <p class="sub">{esc(typ_label)} &middot; {groesse}</p>
  </div>
  <a class="btn" href="{data_url}" download="{esc(pfad.name)}">Herunterladen</a>
</article>'''


def bauen(pfade: list[Path], titel: str) -> str:
    karten = "\n".join(datei_karte(p) for p in pfade)
    return TEMPLATE.format(titel=esc(titel), karten=karten)


TEMPLATE = '''<title>{titel}</title>
<meta name="description" content="Dateien im Original herunterladen - zum Oeffnen, Bearbeiten und Drucken im jeweiligen Programm." />
<style>
  :root {{ --bg:#EFEEEA; --panel:#FFFFFF; --ink:#1B1F22; --muted:#6B6F73; --accent:#0079A7; --accent-ink:#00435C; --line:#D8DBDD; }}
  @media (prefers-color-scheme: dark) {{
    :root:not([data-theme="light"]) {{ --bg:#14171A; --panel:#1D2124; --ink:#EDEFF1; --muted:#9BA0A5; --accent:#4DB8E0; --accent-ink:#7FD1EE; --line:#33383C; }}
  }}
  :root[data-theme="dark"] {{ --bg:#14171A; --panel:#1D2124; --ink:#EDEFF1; --muted:#9BA0A5; --accent:#4DB8E0; --accent-ink:#7FD1EE; --line:#33383C; }}
  * {{ box-sizing:border-box; }}
  body {{ margin:0; padding-inline:16px; padding-block:24px; background:var(--bg); color:var(--ink); font-family:Arial, Helvetica, sans-serif; }}
  .wrap {{ max-width:960px; margin:0 auto; }}
  h1 {{ font-size:clamp(20px,4vw,28px); color:var(--accent-ink); margin:0 0 6px; text-wrap:balance; }}
  .sub-page {{ color:var(--muted); font-size:14px; margin:0 0 24px; max-width:65ch; }}
  .grid {{ display:grid; grid-template-columns:repeat(auto-fill, minmax(240px,1fr)); gap:16px; }}
  .card {{ background:var(--panel); border:1px solid var(--line); border-radius:10px; overflow:hidden; display:flex; flex-direction:column; box-shadow:0 6px 20px -14px rgba(0,0,0,.3); }}
  .preview {{ background:var(--bg); display:flex; align-items:center; justify-content:center; min-height:140px; border-bottom:1px solid var(--line); overflow:hidden; }}
  .img-preview img {{ width:100%; height:140px; object-fit:contain; }}
  .pdf-preview embed {{ width:100%; height:180px; }}
  .no-preview {{ padding:20px; }}
  .pptx-preview {{ padding:10px; align-items:stretch; }}
  .pptx-slide {{ position:relative; width:100%; background:#FFFFFF; border:1px solid var(--line); border-radius:4px; overflow:hidden; color:#1B1F22; font-family:Arial, Helvetica, sans-serif; }}
  .docx-preview {{ align-items:flex-start; padding:14px; height:180px; }}
  .docx-page {{
    width:100%; height:100%; overflow:hidden; position:relative;
    background:#FFFFFF; border:1px solid var(--line); border-radius:4px;
    padding:12px 14px; font-family:Georgia, 'Times New Roman', serif; color:#1B1F22;
  }}
  .docx-page h3 {{ font-size:13px; margin:0 0 6px; }}
  .docx-page p {{ font-size:11px; line-height:1.4; margin:0 0 6px; color:#3A3D40; }}
  .docx-page .fade-hint {{ color:var(--muted); font-style:italic; }}
  .docx-page::after {{
    content:""; position:absolute; left:0; right:0; bottom:0; height:36px;
    background:linear-gradient(to bottom, transparent, #FFFFFF);
  }}
  .file-badge {{ display:inline-block; padding:6px 12px; border-radius:6px; background:var(--accent); color:#FFFFFF; font-size:12px; font-weight:700; letter-spacing:.02em; }}
  .meta {{ padding:12px 14px 4px; flex:1; }}
  .meta h2 {{ font-size:14px; margin:0 0 4px; word-break:break-word; color:var(--ink); }}
  .meta .sub {{ font-size:12px; color:var(--muted); margin:0; }}
  .btn {{ margin:12px 14px 14px; text-align:center; padding:9px 14px; border-radius:6px; background:var(--accent); color:#FFFFFF; font-weight:700; font-size:13px; text-decoration:none; }}
  footer {{ margin-top:28px; padding-top:14px; border-top:1px solid var(--line); font-size:12px; color:var(--muted); }}
</style>
<div class="wrap">
  <header>
    <h1>{titel}</h1>
    <p class="sub-page">Original-Dateien zum Herunterladen. Oeffnen, Bearbeiten und Drucken passiert im jeweiligen Programm (PowerPoint, Word, Bildbetrachter, ...), nicht in dieser Seite.</p>
  </header>
  <div class="grid">{karten}</div>
  <footer>Kienbacher &middot; Datei-Vorschau</footer>
</div>
'''


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("dateien", type=Path, nargs="+")
    parser.add_argument("--out", type=Path, default=Path("vorschau.html"))
    parser.add_argument("--titel", default="Dateien")
    args = parser.parse_args()

    html = bauen(args.dateien, args.titel)
    args.out.write_text(html, encoding="utf-8")
    print(f"Vorschau erzeugt: {args.out} ({len(html)} Bytes, {len(args.dateien)} Dateien)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
